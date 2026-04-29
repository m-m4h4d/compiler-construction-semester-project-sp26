import os
from main import compile_c_code
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # --- Cover Page ---
    # Add some spacing at the top
    for _ in range(5):
        doc.add_paragraph()
        
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Compiler Construction Project Report")
    run.bold = True
    run.font.size = Pt(36)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Spring 2026")
    run.font.size = Pt(24)
    
    for _ in range(10):
        doc.add_paragraph()
        
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run("Modular Compiler for a Subset of C\nLexer, Parser, Semantic Analyzer, TAC & LLVM Backend")
    run.font.size = Pt(14)
    
    doc.add_page_break()
    
    # --- 1. Introduction ---
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph(
        "This document presents the design and implementation of a modular compiler for a subset of the C programming language. "
        "The compiler is implemented entirely in Python and features Lexical Analysis, Syntax Analysis (Recursive Descent Parsing), "
        "Semantic Analysis (Symbol Table for scope checking), Code Optimization (Constant Folding and Dead Code Elimination), "
        "Intermediate Representation (TAC) generation, and finally LLVM IR Backend generation."
    )
    
    # --- 2. Subset Definition & Grammar ---
    doc.add_heading('2. Subset Definition & Grammar', level=1)
    doc.add_heading('Features Supported', level=2)
    features = [
        "Data Types: int, float, double",
        "Expressions: Arithmetic (+, -, *, /, %), Relational (==, !=, <, >, <=, >=), Logical (&&, ||, !)",
        "Control Flow: if-else, while, for loops",
        "Functions: Parameter passing and return statements",
        "Innovation Features: 1D Arrays, Constant Folding, Dead Code Elimination"
    ]
    for feat in features:
        doc.add_paragraph(feat, style='List Bullet')
        
    doc.add_heading('Context-Free Grammar (Subset)', level=2)
    grammar = (
        "Program -> Decl*\n"
        "Decl -> Type ID ( '(' Params ')' Block | '[' NUMBER ']' ';' | [ '=' Expr ] ';' )\n"
        "Block -> '{' Stmt* '}'\n"
        "Stmt -> IfStmt | WhileStmt | ForStmt | ReturnStmt | ExprStmt | Block\n"
        "IfStmt -> 'if' '(' Expr ')' Stmt [ 'else' Stmt ]\n"
        "WhileStmt -> 'while' '(' Expr ')' Stmt\n"
        "ForStmt -> 'for' '(' ExprStmt Expr ';' Expr ')' Stmt\n"
        "ExprStmt -> [ ID '=' Expr | ID '[' Expr ']' '=' Expr ] ';' | Expr ';'\n"
        "Expr -> LogicalOr"
    )
    p = doc.add_paragraph()
    run = p.add_run(grammar)
    run.font.name = 'Courier New'
    
    # --- 3. Design of Compiler Phases ---
    doc.add_heading('3. Design of Compiler Phases', level=1)
    phases = [
        ("Lexical Analyzer (lexer.py)", "Tokenizes the input code using regular expressions to match keywords, numbers, identifiers, and symbols."),
        ("Syntax Analyzer (parser.py)", "Uses Recursive Descent Parsing to build an Abstract Syntax Tree (AST). It incorporates assignments directly into expressions to support standard C for loops."),
        ("Semantic Analyzer (semantic.py)", "Traverses the AST with a stack-based Symbol Table to enforce variable scoping, duplicate declarations, and array access checks."),
        ("Optimizer (optimizer.py)", "Performs Constant Folding on arithmetic expressions and eliminates dead code (e.g. if(0) blocks or while(0))."),
        ("TAC Generator (tac.py)", "A post-order AST traversal to generate flat Three-Address Code with explicit labels and temporaries."),
        ("LLVM Backend (codegen.py)", "Traverses the AST to generate standard LLVM Intermediate Representation (.ll). It handles memory allocation (alloca), loads, and stores to properly reflect local scoping.")
    ]
    for title, desc in phases:
        p = doc.add_paragraph(style='List Number')
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(desc)
        
    # --- 4. Limitations ---
    doc.add_heading('4. Limitations', level=1)
    limits = [
        "Single-dimensional arrays only.",
        "Pointers and structs are not supported.",
        "Missing short-circuiting in logical operators at the LLVM level (currently evaluates both sides).",
        "Only simple typing (implicit promotion between int and float is limited)."
    ]
    for limit in limits:
        doc.add_paragraph(limit, style='List Bullet')
        
    # --- 5. Compilation Instructions ---
    doc.add_heading('5. Compilation Instructions', level=1)
    doc.add_paragraph("The compiler is designed to run in a standard Python 3 environment without external dependencies.")
    instructions = [
        "Place all source files (lexer.py, parser.py, ast_nodes.py, semantic.py, optimizer.py, tac.py, codegen.py, main.py) in a directory.",
        "Run the compiler on a C source file: python main.py input.c",
        "The compiler will produce output.tac (Three-Address Code) and output.ll (LLVM IR).",
        "To compile the LLVM IR to a native executable, use clang: clang output.ll -o output.exe"
    ]
    for inst in instructions:
        doc.add_paragraph(inst, style='List Bullet')
        
    # --- 6. Test Cases & Output ---
    doc.add_heading('6. Test Cases & Output', level=1)
    
    tests = [
        ('test1_arithmetic.c', 'Simple Arithmetic and Constant Folding'),
        ('test2_loops.c', 'Loops and Dead Code Elimination'),
        ('test3_functions_arrays.c', 'Function Calls and Arrays')
    ]
    
    for test_file, desc in tests:
        if not os.path.exists(test_file):
            continue
            
        with open(test_file, 'r') as f:
            code = f.read()
        
        tac, llvm = compile_c_code(code, optimize=True)
        
        doc.add_heading(f"Test Case: {desc} ({test_file})", level=2)
        
        doc.add_paragraph("Source Code:", style='Heading 3')
        p = doc.add_paragraph()
        run = p.add_run(code)
        run.font.name = 'Courier New'
        
        doc.add_paragraph("Generated TAC:", style='Heading 3')
        p = doc.add_paragraph()
        run = p.add_run("\n".join(tac))
        run.font.name = 'Courier New'
        
        doc.add_paragraph("Generated LLVM IR:", style='Heading 3')
        p = doc.add_paragraph()
        run = p.add_run(llvm)
        run.font.name = 'Courier New'
        
        doc.add_page_break()

    # --- 7. Complete Executable Code ---
    doc.add_heading('7. Complete Executable Code', level=1)
    compiler_files = ['ast_nodes.py', 'lexer.py', 'parser.py', 'semantic.py', 'optimizer.py', 'tac.py', 'codegen.py', 'main.py']
    for c_file in compiler_files:
        if not os.path.exists(c_file):
            continue
            
        doc.add_heading(f"{c_file}", level=2)
        with open(c_file, 'r') as f:
            c_code = f.read()
        p = doc.add_paragraph()
        run = p.add_run(c_code)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        
    output_path = '../Compiler_Project_Report.docx'
    doc.save(output_path)
    print(f"Report generated successfully at {output_path}")

if __name__ == "__main__":
    create_report()
